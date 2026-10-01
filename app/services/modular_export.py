"""Native, editable PPTX from accepted modular slide state."""

from __future__ import annotations

from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from app.project_schemas import ProjectResponse


PALETTES = {
    "clean_editorial": ("FFFCF7", "111827", "57534E", "334155", "FFFFFF"),
    "dark_tech_pitch": ("0B1020", "F8FAFC", "A3B2C8", "22C55E", "162238"),
    "infographic_bright": ("F0F9FF", "0F172A", "0369A1", "0EA5E9", "FFFFFF"),
}


def _rgb(hex_code: str) -> RGBColor:
    return RGBColor.from_string(hex_code)


def _rect(slide, x: float, y: float, width: float, height: float, color: str) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    shape.line.fill.background()


def _text(slide, value: str, x: float, y: float, width: float, height: float,
          size: int, color: str, *, bold: bool = False) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    paragraph = frame.paragraphs[0]
    paragraph.text = value
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = _rgb(color)


def render_project_pptx(project: ProjectResponse) -> bytes:
    ordered = sorted(project.outline, key=lambda item: item.order)
    ready = {slide.id: slide for slide in project.slides if slide.status == "ready"}
    if len(ordered) != 5 or len(ready) != 5 or any(item.id not in ready for item in ordered):
        raise ValueError("All five slides must be ready before export")

    background, foreground, muted, accent, panel = PALETTES[project.theme]
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    blank = deck.slide_layouts[6]

    for item in ordered:
        saved = ready[item.id]
        title = saved.blocks.title.text
        body = saved.blocks.body.text
        source = saved.blocks.source_label.text
        slide = deck.slides.add_slide(blank)
        _rect(slide, 0, 0, 13.333, 7.5, background)

        if saved.visual:
            visual = saved.visual
            _text(slide, title, .85, 1.1, 11.7, 1.1, 32, foreground, bold=True)
            _text(slide, body.replace("|", " / "), .85, 2.35, 11.7, 1.25, 21, muted)
            if visual.kind == "process":
                width = 11.65 / len(visual.labels)
                for index, label in enumerate(visual.labels):
                    x = .85 + index * width
                    _rect(slide, x, 4.1, width - .3, 1.75, panel)
                    _text(slide, str(index + 1), x + .15, 4.22, width - .6, .35, 14, accent, bold=True)
                    _text(slide, label, x + .15, 4.72, width - .6, .9, 19, foreground)
            elif visual.kind == "bars":
                maximum = max(visual.values)
                for index, (label, value) in enumerate(zip(visual.labels, visual.values)):
                    y = 3.9 + index * .58
                    _text(slide, label, .85, y, 3.3, .5, 17, foreground)
                    _rect(slide, 4.3, y + .1, value / maximum * 5.5, .28, accent)
                    _text(slide, f"{value:g}{visual.unit}", 10.1, y, 2.25, .5, 17, foreground)
        elif item.layout_type == "title":
            _text(slide, title, 0.85, 1.55, 11.6, 2.2, 42, foreground, bold=True)
            _text(slide, body, 0.85, 4.4, 10.8, 1.4, 23, muted)
        elif item.layout_type == "comparison":
            _text(slide, title, 0.85, 1.1, 11.7, 0.85, 32, foreground, bold=True)
            parts = body.split("|", 1)
            left = parts[0].strip()
            right = parts[1].strip() if len(parts) == 2 else "Second point needed"
            _rect(slide, 0.85, 2.25, 5.65, 3.65, panel)
            _rect(slide, 6.8, 2.25, 5.65, 3.65, panel)
            _text(slide, left, 1.1, 2.65, 5.1, 2.9, 22, foreground)
            _text(slide, right, 7.05, 2.65, 5.1, 2.9, 22, foreground)
        else:
            _text(slide, title, 0.85, 1.05, 11.7, 1.05, 34, foreground, bold=True)
            _rect(slide, 0.85, 2.35, 11.65, 3.55, panel)
            _text(slide, body, 1.15, 2.65, 10.95, 2.95, 25, foreground)

        _rect(slide, 0.85, 6.55, 11.65, 0.02, accent)
        _text(slide, source, 0.85, 6.7, 11.5, 0.45, 12, muted)

    output = BytesIO()
    deck.save(output)
    return output.getvalue()
