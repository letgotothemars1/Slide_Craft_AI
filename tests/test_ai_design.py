import tempfile
from copy import deepcopy
import json
from unittest.mock import MagicMock
from app.services.llm_service import OpenAILLMService
from app.project_schemas import SourceRef
import unittest
from pathlib import Path
from io import BytesIO
from unittest.mock import patch
from pptx import Presentation
from fastapi import BackgroundTasks
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from app import repository
from app.db import Base
from app.project_schemas import DesignPlan, ProjectCreateRequest, OutlineApproveRequest, OutlineRevisionRequest, BlockEditRequest
from app.routers.projects import start_design, retry_design, edit_slide_block, hide_visual, reset_slide_block
from app.services import ai_design, modular_build, modular_recovery
from app.services.modular_build import build_slide_from_outline
from app.services.outline_service import starter_outline
from app.services.modular_export import render_project_pptx


class AIDesignTest(unittest.TestCase):
    def setUp(self):
        self.dir=tempfile.TemporaryDirectory(); self.engine=create_engine(f"sqlite:///{Path(self.dir.name)/'design.db'}")
        Base.metadata.create_all(self.engine); self.sessions=sessionmaker(bind=self.engine)
        self.patches=[patch.object(module,'SessionLocal',self.sessions) for module in (ai_design,modular_build,modular_recovery)]
        self.patches.append(patch('app.routers.projects.get_llm_service',return_value=object()))
        self.patches.append(patch('app.services.design_quality.review_design',return_value={'approved':True,'issues':[],'arrangement':'columns'}))
        self.patches.append(patch('app.services.design_quality.render_scene',return_value=(b'png',[])))
        for p in self.patches: p.start()
        with self.sessions() as session:
            project=repository.create_project(session,ProjectCreateRequest(assignment_text='Five slides',context_pack_text='Thesis: cautious'))
            self.id=project.id; outline=starter_outline('Thesis: cautious')
            project.outline_json=[r.model_dump() for r in outline]
            project.slides_json=[build_slide_from_outline(r).model_dump() for r in outline]
            project.phase='outline_draft'; session.commit()
    def tearDown(self):
        for p in reversed(self.patches): p.stop()
        self.engine.dispose(); self.dir.cleanup()
    def state(self):
        with self.sessions() as session: return repository.project_response(session,repository.get_project(session,self.id))
    def start(self):
        with self.sessions() as session: return start_design(self.id,OutlineApproveRequest(expected_revision=self.state().revision,theme='clean_editorial'),BackgroundTasks(),session)
    def plan(self,project,item,slide,previous):
        return DesignPlan(layout='hero' if item.order==1 else 'comparison' if '|' in slide['blocks']['body']['text'] else 'editorial',emphasis='quiet',rationale='A clear hierarchy.')
    def test_design_preserves_all_text_and_exports_shared_geometry(self):
        before=self.state(); self.start()
        def generate(*args):
            current=self.state()
            self.assertEqual(current.phase,'designing')
            self.assertEqual(sum(r.design_status=='generating' for r in current.slides),1)
            return self.plan(*args)
        with patch.object(ai_design,'generate_design',side_effect=generate): ai_design.run_design(self.id)
        after=self.state(); self.assertEqual(after.phase,'ready')
        self.assertEqual([r.blocks for r in after.slides],[r.blocks for r in before.slides])
        self.assertTrue(all(r.design_status=='ready' and r.scene for r in after.slides))
        deck=Presentation(BytesIO(render_project_pptx(after)))
        for source,slide in zip(after.slides,deck.slides):
            self.assertEqual(len(source.scene),len(slide.shapes))
            for element,shape in zip(source.scene,slide.shapes):
                self.assertAlmostEqual(shape.left/914400,element.x*13.333/100,places=4)
                self.assertAlmostEqual(shape.top/914400,element.y*13.333/100,places=4)
                self.assertGreaterEqual(element.x,0); self.assertLessEqual(element.x+element.w,100.01)
                self.assertLessEqual(element.y+element.h,56.26)
                if element.kind=='text': self.assertEqual(shape.text.replace('\v','\n'),element.text)
            self.assertIn(source.blocks.source_label.text,slide.notes_slide.notes_text_frame.text)
    def test_failure_and_retry_keep_completed_designs(self):
        self.start()
        def generate(project,item,slide,previous):
            if item.order==3: raise ValueError('Deliberate failure')
            return self.plan(project,item,slide,previous)
        with patch.object(ai_design,'generate_design',side_effect=generate),self.assertLogs(ai_design.logger,level='ERROR'): ai_design.run_design(self.id)
        failed=self.state(); self.assertEqual(failed.phase,'outline_draft'); self.assertEqual(failed.slides[2].design_status,'error')
        with self.sessions() as session: retry_design(self.id,'s3',OutlineRevisionRequest(expected_revision=failed.revision),BackgroundTasks(),session)
        with patch.object(ai_design,'generate_design',side_effect=self.plan) as generate: ai_design.run_design(self.id,only_slide_id='s3')
        self.assertEqual(generate.call_count,1); after=self.state(); self.assertEqual(after.phase,'ready')
        for index in (0,1,3,4): self.assertEqual(after.slides[index],failed.slides[index])
    def test_concurrent_edit_discards_stale_design(self):
        self.start()
        def generate(project,item,slide,previous):
            if item.order==1:
                with self.sessions() as session: edit_slide_block(self.id,'s1','title',BlockEditRequest(expected_revision=self.state().revision,text='Accepted edit'),session)
            return self.plan(project,item,slide,previous)
        with patch.object(ai_design,'generate_design',side_effect=generate): ai_design.run_design(self.id)
        state=self.state(); self.assertEqual(state.slides[0].blocks.title.text,'Accepted edit'); self.assertIsNone(state.slides[0].design)
        self.assertEqual(state.slides[0].design_status,'error')
    def test_restart_makes_design_retryable_without_losing_content(self):
        before=self.state(); self.start(); modular_recovery.recover_interrupted_work(); after=self.state()
        self.assertEqual(after.phase,'outline_draft'); self.assertTrue(all(r.design_status=='error' for r in after.slides))
        self.assertEqual([r.blocks for r in before.slides],[r.blocks for r in after.slides])

    def test_provider_schema_excludes_charts_without_accepted_numbers(self):
        service=OpenAILLMService.__new__(OpenAILLMService); service.model='test'; service.client=MagicMock()
        service._extract_output_text=lambda _:json.dumps({'layout':'editorial','emphasis':'quiet','visual':{'kind':'none','labels':[],'values':[],'unit':''},'rationale':'Emphasize the cautious claim.'})
        with self.sessions() as session:
            project=repository.get_project(session,self.id); item=starter_outline('Thesis: cautious')[1]
            with patch.object(ai_design,'get_llm_service',return_value=service):
                plan=ai_design.generate_design(project,item,project.slides_json[1],[])
        schema=service.client.responses.create.call_args.kwargs['text']['format']['schema']
        self.assertNotIn('chart',schema['properties']['layout']['enum'])
        self.assertNotIn('bars',schema['properties']['visual']['properties']['kind']['enum'])
        self.assertEqual(plan.layout,'editorial')

    def test_provider_cannot_add_pdf_numbers_absent_from_accepted_body(self):
        service=OpenAILLMService.__new__(OpenAILLMService); service.model='test'; service.client=MagicMock()
        service._extract_output_text=lambda _:json.dumps({'layout':'chart','emphasis':'quiet','visual':{'kind':'bars','labels':['Before','After'],'values':[42,29],'unit':'%'},'rationale':'Comparison.'})
        with self.sessions() as session:
            project=repository.get_project(session,self.id); item=starter_outline('Thesis: cautious')[1]
            item.suggested_refs=[SourceRef(document_id='d',filename='demo.pdf',page_number=1,excerpt='Synthetic 42% and 29%.')]
            with patch.object(ai_design,'get_llm_service',return_value=service),self.assertRaisesRegex(ValueError,'accepted numbers'):
                ai_design.generate_design(project,item,project.slides_json[1],[])

    def test_hide_visual_preserves_final_design_and_accepted_body(self):
        self.start()
        with patch.object(ai_design,'generate_design',side_effect=self.plan): ai_design.run_design(self.id)
        with self.sessions() as session:
            project=repository.get_project(session,self.id)
            slides=deepcopy(project.slides_json)
            slides[0]['design']['layout']='process'
            slides[0]['design']['visual']={'kind':'process','labels':['Review','Decide'],'values':[],'unit':''}
            project.slides_json=slides; session.commit()
            before=self.state()
            after=hide_visual(self.id,'s1',OutlineRevisionRequest(expected_revision=before.revision),session)
        self.assertEqual(after.phase,'ready')
        self.assertEqual(after.slides[0].design.layout,'hero')
        self.assertIsNone(after.slides[0].design.visual)
        self.assertEqual(after.slides[0].blocks,before.slides[0].blocks)
        self.assertTrue(after.slides[0].scene)

    def test_reset_from_outline_invalidates_final_composition(self):
        self.start()
        with patch.object(ai_design,'generate_design',side_effect=self.plan): ai_design.run_design(self.id)
        with self.sessions() as session:
            after=reset_slide_block(self.id,'s1','body',OutlineRevisionRequest(expected_revision=self.state().revision),session)
        self.assertIsNone(after.slides[0].design)
        self.assertEqual(after.slides[0].design_status,'none')
        with self.assertRaisesRegex(ValueError,'outdated'): render_project_pptx(after)

    def test_theme_inversion_retains_text_and_visual_contrast(self):
        from app.services.design_scene import build_scene, _contrast
        for theme in ('clean_editorial','dark_tech_pitch','infographic_bright'):
            for emphasis in ('quiet','inverse'):
                slide=self.state().slides[0]
                slide.design=DesignPlan(layout='hero',emphasis=emphasis,rationale='Contrast test')
                scene=build_scene(slide,theme,1)
                for element in (r for r in scene if r.kind=='text'):
                    surface=next(r for r in reversed(scene[:scene.index(element)]) if r.kind=='rect' and r.x<=element.x and r.y<=element.y and r.x+r.w>=element.x+element.w and r.y+r.h>=element.y+element.h)
                    self.assertGreaterEqual(_contrast(element.color,surface.color),4.5,(theme,emphasis,element.text))

    def test_design_preserves_reviewed_visuals_and_theme_including_old_plans(self):
        from app.project_schemas import DraftVisual
        from app.services.design_constraints import preserve_reviewed_structure, allowed_layouts
        visual=DraftVisual(kind='process',labels=['Flag uncertain items','Staff review'],values=[],unit='')
        proposed=DesignPlan(layout='editorial',emphasis='inverse',rationale='Old model plan.')
        locked=preserve_reviewed_structure(proposed,'content',visual)
        self.assertEqual(locked.layout,'process')
        self.assertEqual(locked.visual,visual)
        self.assertEqual(locked.emphasis,'quiet')
        self.assertEqual(allowed_layouts('comparison',None),['comparison'])
        self.assertEqual(allowed_layouts('title',None),['hero'])
        self.start()
        with patch.object(ai_design,'generate_design',side_effect=self.plan): ai_design.run_design(self.id)
        with self.sessions() as session:
            project=repository.get_project(session,self.id); slides=deepcopy(project.slides_json)
            slides[0]['design']['emphasis']='inverse'
            slides[1]['visual']=visual.model_dump()
            slides[1]['design']['layout']='editorial'; slides[1]['design']['visual']=None
            project.slides_json=slides; session.commit()
        after=self.state()
        self.assertEqual(after.slides[0].design.emphasis,'quiet')
        self.assertEqual(after.slides[1].design.layout,'process')
        self.assertEqual(after.slides[1].design.visual,visual)
        self.assertEqual(len({slide.scene[0].color for slide in after.slides}),1)
        deck=Presentation(BytesIO(render_project_pptx(after)))
        self.assertIn('Staff review',' '.join(shape.text for shape in deck.slides[1].shapes if shape.has_text_frame))

    def test_quality_refinement_is_bounded_and_preserves_content(self):
        before=self.state(); self.start(); calls=[]
        def compose(project,item,slide,previous,feedback=None):
            calls.append((item.id,feedback)); return self.plan(project,item,slide,previous)
        reviews=[{'approved':False,'issues':['Crowded columns'],'arrangement':'rows'}]+[{'approved':True,'issues':[],'arrangement':'rows'}]*5
        with patch.object(ai_design,'generate_design',side_effect=compose),patch('app.services.design_quality.review_design',side_effect=reviews): ai_design.run_design(self.id)
        after=self.state()
        self.assertEqual(after.phase,'ready'); self.assertEqual(after.slides[0].quality_attempts,2)
        self.assertEqual(len(calls),6); self.assertEqual(calls[1][1]['issues'],['Crowded columns'])
        self.assertEqual(after.slides[0].design.arrangement,'rows')
        self.assertEqual([r.blocks for r in after.slides],[r.blocks for r in before.slides])

    def test_quality_failure_keeps_candidate_and_blocks_export(self):
        self.start()
        with patch.object(ai_design,'generate_design',side_effect=lambda *args,**kwargs:self.plan(*args)),patch('app.services.design_quality.review_design',return_value={'approved':False,'issues':['Unreadable text'],'arrangement':'rows'}): ai_design.run_design(self.id)
        after=self.state(); self.assertEqual(after.phase,'outline_draft')
        self.assertTrue(all(r.design_status=='error' and r.quality_attempts==3 and r.design for r in after.slides))
        with self.assertRaises(ValueError):render_project_pptx(after)

    def test_section_edit_keeps_data_bindings_and_invalidates_only_its_design(self):
        from app.project_schemas import SectionEditRequest, SlideSection
        from app.routers.projects import edit_sections
        from app.services.semantic_sections import validated_sections
        sections=validated_sections([{'heading':'Staff review','text':'Staff review takes 18 minutes.'},{'heading':'Maintenance','text':'Maintenance takes two hours.'}], 'Staff review takes 18 minutes. Maintenance takes two hours.','s2')
        with self.assertRaisesRegex(ValueError,'every accepted word'):validated_sections([{'heading':'Other','text':'Staff review takes 20 minutes.'}],'Staff review takes 18 minutes.','s2')
        self.start()
        with patch.object(ai_design,'generate_design',side_effect=self.plan):ai_design.run_design(self.id)
        before=self.state()
        with self.sessions() as session:after=edit_sections(self.id,'s2',SectionEditRequest(expected_revision=before.revision,sections=sections),session)
        self.assertEqual(after.slides[1].sections,sections)
        self.assertIsNone(after.slides[1].design)
        self.assertEqual(after.slides[0].design,before.slides[0].design)
        from app.services.design_scene import build_scene
        from app.services.design_quality import render_scene
        slide=after.slides[1].model_copy(update={'design':DesignPlan(layout='editorial',emphasis='quiet',rationale='Keep sections',arrangement='rows')})
        scene=build_scene(slide,after.theme,2)
        self.assertEqual({e.section_id for e in scene if e.section_id},{s.id for s in sections})
        self.assertEqual([e.text for e in scene if e.section_field=='text'],[s.text for s in sections])
