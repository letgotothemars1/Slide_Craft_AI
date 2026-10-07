"""Provider-free integration checks for the staged tool executor and budget."""
import json
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from fastapi import BackgroundTasks
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from app import repository
from app.db import Base
from app.project_schemas import ProjectCreateRequest, BuildRequest, OutlineApproveRequest
from app.routers.projects import start_draft, start_design
from app.services import live_draft, ai_design, modular_build, modular_recovery
from app.services.deck_design import generate_deck_design, validate_plan
from app.services.llm_service import OpenAILLMService
from app.services.presentation_workflow import BudgetExhausted, operation, record_usage
from app.services.modular_export import render_project_pptx
from pptx import Presentation
from io import BytesIO

class StagedWorkflowTest(unittest.TestCase):
    def setUp(self):
        self.directory=tempfile.TemporaryDirectory()
        self.engine=create_engine(f"sqlite:///{Path(self.directory.name)/'workflow.db'}")
        Base.metadata.create_all(self.engine)
        self.sessions=sessionmaker(bind=self.engine,expire_on_commit=False)
        self.patches=[patch.object(m,'SessionLocal',self.sessions) for m in (live_draft,ai_design,modular_build,modular_recovery)]
        self.preflight_patch=patch.object(ai_design,'ensure_deck_variants',create=True)
        self.patches.append(self.preflight_patch)
        self.patches.append(patch('app.services.composition_choices.ensure_composition_variants'))
        for p in self.patches:p.start()
        with self.sessions() as s:
            p=repository.create_project(s,ProjectCreateRequest(assignment_text='Explain a synthetic study in five slides.',context_pack_text='Thesis: review then decide.'))
            self.id=p.id
    def tearDown(self):
        for p in reversed(self.patches):p.stop()
        self.engine.dispose();self.directory.cleanup()
    def state(self):
        with self.sessions() as s:return repository.project_response(s,repository.get_project(s,self.id))
    def begin(self,mode='template'):
        with self.sessions() as s:return start_draft(self.id,BuildRequest(expected_revision=self.state().revision,mode=mode),BackgroundTasks(),s)
    def begin_design(self):
        with self.sessions() as s:return start_design(self.id,OutlineApproveRequest(expected_revision=self.state().revision,theme='clean_editorial'),BackgroundTasks(),s)
    def test_legacy_template_feature_without_sections_can_be_planned(self):
        self.begin();live_draft.run_live_draft(self.id)
        with self.sessions() as session:
            project=repository.get_project(session,self.id)
            rows=deepcopy(project.slides_json)
            rows[0].update(sections=[],composition_preference='feature')
            project.slides_json=rows
            session.commit()
            plan=generate_deck_design(project)
        first=next(a for a in plan['actions'] if a['slide_id']==rows[0]['id'])
        self.assertEqual(first['composition'],'feature')
        self.assertEqual(first['focal_section_id'],'')
    def fake_provider(self):
        service=OpenAILLMService.__new__(OpenAILLMService);service.model='mock';service.client=MagicMock()
        service._extract_output_text=lambda response:response.output_text
        def create(**kw):
            schema=kw['text']['format']['name']
            if schema=='academic_outline':
                data={'slides':[{'title':f'Slide {i}','purpose':'Explain a part','key_message':'Review and decide.','layout_type':'title' if i==1 else 'content'} for i in range(1,6)]}
            elif schema=='academic_slide_body':
                data={'body':'Review the evidence. Decide next steps.', 'speaker_notes':'Explain how evidence review informs the decision and the next steps.', 'visual':{'kind':'none','labels':[],'values':[],'unit':''},
                    'sections':[{'heading':'Review','text':'Review the evidence.'},{'heading':'Decide','text':'Decide next steps.'}]}
                raw=json.dumps(data)
                return iter([SimpleNamespace(type='response.output_text.delta',delta=raw[:45]),SimpleNamespace(type='response.output_text.delta',delta=raw[45:]),
                    SimpleNamespace(type='response.completed',response=SimpleNamespace(usage=SimpleNamespace(input_tokens=10,output_tokens=5)))])
            elif schema=='slide_notes_coherence':
                data={'approved':True,'summary':'Slides and notes support the same argument.','findings':[]}
            elif schema=='slide_composition_variants':
                user=json.loads(kw['input'][1]['content'])
                data={'variants':[{'id':identifier,'label':label,'description':'A different reading order of the accepted argument.',
                    'design':{'layout':user['allowed_layouts'][0],'composition':composition,'arrangement':'rows',
                        'focal_section_id':'','emphasis':'quiet','rationale':'Make the argument legible.'}}
                    for identifier,label,composition in [('selected','Side by side','balanced'),('alternative','Read in sequence','bands'),('out_of_box','Emphasize the question','poster')]]}
            elif schema=='presentation_story_analysis':
                user=json.loads(kw['input'][1]['content'])
                data={'audience_goal':'Help staff decide next steps.','narrative':'Review then decide.','slides':[{'slide_id':r['slide_id'],'communication_goal':r['title'],'visible_priority':r['body'],'spoken_explanation':r.get('speaker_notes',''),'visual_strategy':'Distinguish review and decision.','content_gaps':[]} for r in user['deck']]}
            elif schema=='slide_art_direction':
                user=json.loads(kw['input'][1]['content'].split('\nAllowed numeric')[0])
                feedback=json.loads(kw['input'][1]['content'].split('RENDERED REVIEW FEEDBACK: ')[1])
                data=feedback['current_design']
                if data.get('visual') is None:data['visual']={'kind':'none','labels':[],'values':[],'unit':''}
            elif schema=='complete_deck_review':data={'approved':True,'slides':[]}
            elif schema=='deck_composition_actions':
                user=json.loads(kw['input'][1]['content']);actions=[]
                for i,row in enumerate(user['deck']):
                    composition=row.get('composition_preference') or (('balanced','feature','bands')[i%3] if row['sections'] else 'balanced')
                    actions.append({'action':'apply_composition','slide_id':row['slide_id'],'composition':composition,
                        'focal_section_id':row['sections'][0]['id'] if composition=='feature' else '',
                        'arrangement':'rows','layout':row['allowed_layouts'][0],'rationale':'Different views of accepted information.'})
                data={'direction':'Quiet light editorial hierarchy.', 'actions':actions}
            elif schema=='rendered_design_review':data={'approved':True,'issues':[],'arrangement':'rows'}
            else:raise AssertionError('Unexpected extra model call: '+schema)
            return SimpleNamespace(output_text=json.dumps(data),usage=SimpleNamespace(input_tokens=10,output_tokens=5))
        service.client.responses.create.side_effect=create
        return service
    def test_live_generated_three_variants_and_selected_main_survive_final_design(self):
        service=self.fake_provider()
        automatic_variants=self.patches[-1]
        automatic_variants.stop();self.preflight_patch.stop()
        try:
            modules=['app.routers.projects','app.services.project_outline_llm','app.services.modular_slide_llm','app.services.deck_design','app.services.design_quality','app.services.ai_design']
            with __import__('contextlib').ExitStack() as stack:
                stack.enter_context(patch('app.db.SessionLocal',self.sessions))
                for module in modules:stack.enter_context(patch(module+'.get_llm_service',return_value=service))
                self.begin('model');live_draft.run_live_draft(self.id)
                draft=self.state()
                self.assertTrue(all(len(s.composition_variants)==3 and s.selected_variant_id=='selected' for s in draft.slides))
                self.assertEqual(draft.workflow.model_calls,11)
                with self.sessions() as session:
                    project=repository.get_project(session,self.id);rows=deepcopy(project.slides_json)
                    rows[0]['selected_variant_id']='out_of_box';project.slides_json=rows;session.commit()
                self.begin_design();ai_design.run_design(self.id)
            result=self.state();self.assertEqual(result.phase,'ready')
            self.assertEqual(result.slides[0].design.composition,'poster')
            self.assertEqual(result.slides[0].selected_variant_id,'out_of_box')
            chosen=next(v for v in result.slides[0].composition_variants if v.id=='out_of_box')
            self.assertEqual(chosen.design,result.slides[0].design)
            self.assertEqual([s.blocks for s in draft.slides],[s.blocks for s in result.slides])
            names=[c.kwargs['text']['format']['name'] for c in service.client.responses.create.call_args_list]
            self.assertEqual(names.count('slide_composition_variants'),5)
            self.assertEqual(result.workflow.model_calls,30)
            self.assertEqual(service.client.responses.create.call_count,30)
            self.assertEqual((result.workflow.input_tokens,result.workflow.output_tokens),(300,150))
            self.assertNotIn('composition_candidate_review',names)
        finally:
            automatic_variants.start();self.preflight_patch.start()
    def test_bounded_worst_case_quality_workflow_fits_new_default_budget(self):
        service=self.fake_provider();original=service.client.responses.create.side_effect
        review_count=0;deck_count=0
        def worst_case(**kw):
            nonlocal review_count,deck_count
            name=kw['text']['format']['name']
            if name=='rendered_design_review':
                review_count+=1
                approved=review_count>15 or review_count%3==0
                data={'approved':approved,'issues':[] if approved else ['Strengthen visible hierarchy.'],'arrangement':'rows'}
            elif name=='complete_deck_review':
                deck_count+=1
                data={'approved':deck_count==2,'slides':[] if deck_count==2 else [{'slide_id':f's{i}','issues':['Align the hierarchy with the rest of the deck.']} for i in range(1,6)]}
            else:return original(**kw)
            return SimpleNamespace(output_text=json.dumps(data),usage=SimpleNamespace(input_tokens=10,output_tokens=5))
        service.client.responses.create.side_effect=worst_case
        modules=['app.routers.projects','app.services.project_outline_llm','app.services.modular_slide_llm','app.services.deck_design','app.services.design_quality','app.services.ai_design']
        with __import__('contextlib').ExitStack() as stack:
            for module in modules:stack.enter_context(patch(module+'.get_llm_service',return_value=service))
            self.begin('model');live_draft.run_live_draft(self.id);self.begin_design();ai_design.run_design(self.id)
        result=self.state();self.assertEqual(result.phase,'ready')
        self.assertEqual(result.workflow.request_budget,60)
        self.assertEqual(result.workflow.model_calls,46);self.assertEqual(service.client.responses.create.call_count,46)
        self.assertEqual(review_count,20);self.assertEqual(deck_count,2)
        self.assertTrue(all(slide.quality_attempts==4 for slide in result.slides))

    def test_missing_current_model_variants_block_final_design_without_fallback(self):
        service=self.fake_provider();self.begin();live_draft.run_live_draft(self.id)
        with self.sessions() as session:
            project=repository.get_project(session,self.id);project.build_mode='model';session.commit()
        before=self.state();self.preflight_patch.stop()
        try:
            with patch('app.routers.projects.get_llm_service',return_value=service):self.begin_design()
            with patch('app.services.deck_design.get_llm_service',return_value=service), patch.object(ai_design,'initial_design',side_effect=AssertionError('No fallback when model alternatives are unavailable')) as initial, self.assertLogs(ai_design.logger,level='ERROR'):
                ai_design.run_design(self.id)
            result=self.state();self.assertEqual(result.phase,'outline_draft')
            self.assertTrue(all(s.design_status=='error' for s in result.slides))
            self.assertEqual([s.blocks for s in before.slides],[s.blocks for s in result.slides])
            self.assertEqual([s.speaker_notes for s in before.slides],[s.speaker_notes for s in result.slides])
            initial.assert_not_called()
        finally:
            self.preflight_patch.start()

    def test_legacy_model_draft_generates_exactly_three_before_final_styling(self):
        service=self.fake_provider();self.begin();live_draft.run_live_draft(self.id)
        with self.sessions() as session:
            # A ready key-free cache must not masquerade as AI alternatives when switching to model mode.
            from app.services.composition_choices import generate_composition_variants,content_fingerprint
            from app.project_schemas import ProjectSlide,OutlineItem
            project=repository.get_project(session,self.id);rows=deepcopy(project.slides_json)
            for row in rows:
                item=next(OutlineItem.model_validate(i) for i in project.outline_json if i['id']==row['id'])
                row.update(composition_variants=generate_composition_variants(project,item,ProjectSlide.model_validate(row)),
                    selected_variant_id='selected',variants_status='ready',variants_origin='template',
                    variants_fingerprint=content_fingerprint(row,item))
            project.slides_json=rows;project.build_mode='model';session.commit()
        before=self.state();automatic_variants=self.patches[-1]
        automatic_variants.stop();self.preflight_patch.stop()
        try:
            with __import__('contextlib').ExitStack() as stack:
                stack.enter_context(patch('app.db.SessionLocal',self.sessions))
                for module in ['app.routers.projects','app.services.deck_design','app.services.design_quality','app.services.ai_design']:
                    stack.enter_context(patch(module+'.get_llm_service',return_value=service))
                self.begin_design();ai_design.run_design(self.id)
            result=self.state();self.assertEqual(result.phase,'ready')
            self.assertTrue(all(s.design_status=='ready' and s.variants_origin=='model' and len(s.composition_variants)==3 for s in result.slides))
            self.assertEqual([s.blocks for s in before.slides],[s.blocks for s in result.slides])
            names=[c.kwargs['text']['format']['name'] for c in service.client.responses.create.call_args_list]
            self.assertLess(names.index('slide_notes_coherence'),names.index('slide_composition_variants'))
            self.assertEqual(names.count('slide_composition_variants'),5)
            self.assertEqual(result.workflow.model_calls,24)
            self.assertEqual(service.client.responses.create.call_count,24)
            self.assertEqual((result.workflow.input_tokens,result.workflow.output_tokens),(240,120))
            self.assertNotIn('composition_candidate_review',names)
        finally:
            automatic_variants.start();self.preflight_patch.start()

    def test_new_default_does_not_raise_existing_stored_budget(self):
        from app.project_schemas import WorkflowState
        self.assertEqual(WorkflowState().request_budget,60)
        self.assertEqual(WorkflowState.model_validate({'request_budget':40}).request_budget,40)

    def test_verified_reasoning_model_only_uses_high_for_quality_calls(self):
        service=self.fake_provider();service.model='gpt-5.4-mini-2026-03-17'
        modules=['app.routers.projects','app.services.project_outline_llm','app.services.modular_slide_llm','app.services.deck_design','app.services.design_quality','app.services.ai_design']
        with __import__('contextlib').ExitStack() as stack:
            for module in modules:stack.enter_context(patch(module+'.get_llm_service',return_value=service))
            self.begin('model');live_draft.run_live_draft(self.id);self.begin_design();ai_design.run_design(self.id)
        self.assertEqual(self.state().phase,'ready')
        for call in service.client.responses.create.call_args_list:
            kw=call.kwargs;name=kw['text']['format']['name']
            if name in {'academic_outline','academic_slide_body'}:
                self.assertNotIn('reasoning',kw)
            else:
                self.assertEqual(kw['reasoning'],{'effort':'high'},name)
                self.assertNotIn('temperature',kw,name)

    def test_model_design_uses_notes_and_keeps_selected_poster(self):
        service=self.fake_provider()
        self.begin();live_draft.run_live_draft(self.id)
        with self.sessions() as session:
            project=repository.get_project(session,self.id);rows=deepcopy(project.slides_json)
            rows[0]['speaker_notes']='Explain why a human reviews the evidence before deciding.'
            rows[0]['composition_preference']='poster';project.slides_json=rows;project.build_mode='model';session.commit()
        before=self.state()
        with patch('app.routers.projects.get_llm_service',return_value=service):self.begin_design()
        with patch('app.services.deck_design.get_llm_service',return_value=service),patch('app.services.design_quality.get_llm_service',return_value=service),patch('app.services.ai_design.get_llm_service',return_value=service):
            ai_design.run_design(self.id)
        result=self.state();self.assertEqual(result.phase,'ready')
        self.assertEqual(result.slides[0].design.composition,'poster')
        self.assertEqual([s.speaker_notes for s in before.slides],[s.speaker_notes for s in result.slides])
        self.assertEqual([s.sections for s in before.slides],[s.sections for s in result.slides])
        calls=service.client.responses.create.call_args_list
        schemas=[c.kwargs['text']['format']['name'] for c in calls]
        self.assertEqual(schemas[0],'slide_notes_coherence')
        self.assertEqual(schemas[-1],'complete_deck_review')
        for name in ('presentation_story_analysis','deck_composition_actions','rendered_design_review','slide_art_direction','complete_deck_review'):
            selected=next(c.kwargs for c in calls if c.kwargs['text']['format']['name']==name)
            self.assertIn(before.slides[0].speaker_notes,json.dumps(selected['input']))
        self.assertEqual(result.workflow.story_analysis['slides'][0]['spoken_explanation'],before.slides[0].speaker_notes)

    def test_final_deck_rejection_does_not_misreport_ready(self):
        service=self.fake_provider();original=service.client.responses.create.side_effect
        def reviewed(**kw):
            if kw['text']['format']['name']=='complete_deck_review':
                return SimpleNamespace(output_text=json.dumps({'approved':False,'slides':[{'slide_id':'s1','issues':['Accepted evidence does not support the conclusion; styling cannot repair this.']}]}),usage=SimpleNamespace(input_tokens=10,output_tokens=5))
            return original(**kw)
        service.client.responses.create.side_effect=reviewed
        self.begin();live_draft.run_live_draft(self.id)
        with self.sessions() as session:
            project=repository.get_project(session,self.id);project.build_mode='model';session.commit()
        before=self.state()
        with patch('app.routers.projects.get_llm_service',return_value=service):self.begin_design()
        with patch('app.services.deck_design.get_llm_service',return_value=service),patch('app.services.design_quality.get_llm_service',return_value=service),patch('app.services.ai_design.get_llm_service',return_value=service):ai_design.run_design(self.id)
        result=self.state();self.assertEqual(result.phase,'outline_draft');self.assertEqual(result.workflow.stage,'error')
        self.assertEqual(result.slides[0].design_status,'error');self.assertFalse(result.workflow.deck_review['approved'])
        self.assertEqual([s.blocks for s in before.slides],[s.blocks for s in result.slides])
        self.assertTrue(all(s.design is not None for s in result.slides))

    def test_exhaustion_before_global_review_preserves_designs_but_not_ready(self):
        service=self.fake_provider();self.begin();live_draft.run_live_draft(self.id)
        with self.sessions() as session:
            project=repository.get_project(session,self.id);project.build_mode='model'
            project.workflow_json={**project.workflow_json,'request_budget':18};session.commit()
        with patch('app.routers.projects.get_llm_service',return_value=service):self.begin_design()
        with patch('app.services.deck_design.get_llm_service',return_value=service),patch('app.services.design_quality.get_llm_service',return_value=service),patch('app.services.ai_design.get_llm_service',return_value=service):ai_design.run_design(self.id)
        result=self.state();self.assertEqual(result.phase,'outline_draft');self.assertEqual(result.workflow.stage,'budget_exhausted')
        self.assertEqual(result.workflow.model_calls,18);self.assertTrue(all(s.design is not None for s in result.slides))
        self.assertTrue(all(s.design_status=='error' for s in result.slides))

    def test_keyfree_workflow_does_not_initialize_provider(self):
        with patch('app.routers.projects.get_llm_service',side_effect=AssertionError('No keyfree provider calls')),patch('app.services.deck_design.get_llm_service',side_effect=AssertionError('No keyfree planning calls')),patch('app.services.design_quality.get_llm_service',side_effect=AssertionError('No keyfree vision calls')):
            self.begin();live_draft.run_live_draft(self.id);self.begin_design();ai_design.run_design(self.id)
        result=self.state();self.assertEqual(result.phase,'ready');self.assertEqual(result.workflow.model_calls,0)
        self.assertEqual(len(result.workflow.deck_design['actions']),5)
        self.assertEqual(len(Presentation(BytesIO(render_project_pptx(result))).slides),5)
    def test_budget_stops_before_next_provider_and_keeps_completed_content(self):
        service=self.fake_provider()
        with self.sessions() as s:
            p=repository.get_project(s,self.id);p.workflow_json={'request_budget':2};s.commit()
        with patch('app.routers.projects.get_llm_service',return_value=service),patch('app.services.project_outline_llm.get_llm_service',return_value=service),patch('app.services.modular_slide_llm.get_llm_service',return_value=service):
            self.begin('model');live_draft.run_live_draft(self.id)
        result=self.state();self.assertEqual(result.workflow.stage,'budget_exhausted')
        self.assertEqual(result.workflow.model_calls,2);self.assertEqual(service.client.responses.create.call_count,2)
        self.assertEqual(result.slides[0].status,'ready');self.assertTrue(result.slides[0].blocks.body.text)
        self.assertTrue(all(s.status=='error' for s in result.slides[1:]));self.assertEqual(result.phase,'outline_draft')
    def test_deck_plan_rejects_unknown_duplicate_and_foreign_focal_ids(self):
        self.begin();live_draft.run_live_draft(self.id)
        with self.sessions() as s:
            p=repository.get_project(s,self.id);valid=generate_deck_design(p)
            for mutate in (lambda d:d['actions'][0].update(slide_id='alien'),lambda d:d['actions'][1].update(slide_id='s1'),lambda d:d['actions'][0].update(focal_section_id='alien-section'),lambda d:d['actions'][1].update(layout='chart')):
                data=deepcopy(valid);mutate(data)
                with self.assertRaises(ValueError):validate_plan(data,p)
    def test_concurrent_content_edit_during_planning_discards_all_actions(self):
        self.begin();live_draft.run_live_draft(self.id);self.begin_design()
        original=generate_deck_design
        def changed(project):
            result=original(project)
            def mutate(current):
                rows=deepcopy(current.slides_json);rows[0]['blocks']['title']['text']='Accepted newer title';rows[0]['blocks']['title']['revision']+=1
                return {'slides_json':rows}
            modular_build._mutate(self.id,mutate);return result
        with patch('app.services.deck_design.generate_deck_design',side_effect=changed),self.assertLogs(ai_design.logger,level='ERROR'):ai_design.run_design(self.id)
        result=self.state();self.assertEqual(result.slides[0].blocks.title.text,'Accepted newer title')
        self.assertTrue(all(s.design is None and s.design_status=='error' for s in result.slides))
        self.assertIsNone(result.workflow.deck_design)
    def test_recovery_closes_running_events_and_keeps_budget(self):
        self.begin()
        with operation(self.id,'plan_deck',model=True):
            modular_recovery.recover_interrupted_work()
        result=self.state();self.assertEqual(result.phase,'intake');self.assertEqual(result.workflow.model_calls,1)
        self.assertEqual(result.workflow.stage,'error')
    def test_manual_revision_budget_is_visible_and_retains_accepted_text(self):
        from app.services import modular_regenerate
        self.begin();live_draft.run_live_draft(self.id)
        before=self.state().slides[0].blocks.body.text
        with self.sessions() as session:
            project=repository.get_project(session,self.id)
            rows=deepcopy(project.slides_json)
            block=rows[0]['blocks']['body'];block['status']='generating';block['revision']+=1
            token=block['revision'];project.slides_json=rows
            project.workflow_json={**project.workflow_json,'model_calls':40,'request_budget':40}
            session.commit()
        with patch.object(modular_regenerate,'SessionLocal',self.sessions),patch.object(modular_regenerate,'generate_block_text',side_effect=AssertionError('No provider at exhausted budget')),self.assertLogs(modular_regenerate.logger,level='ERROR'):
            modular_regenerate.run_regeneration(self.id,'s1','body',token)
        result=self.state()
        self.assertEqual(result.workflow.stage,'budget_exhausted')
        self.assertEqual(result.slides[0].blocks.body.text,before)
        self.assertEqual(result.slides[0].blocks.body.status,'error')
        self.assertIn('budget',result.slides[0].blocks.body.error.lower())
        self.assertEqual(result.workflow.model_calls,40)

    def test_atomic_budget_rejects_without_increment(self):
        with self.sessions() as s:
            p=repository.get_project(s,self.id);p.workflow_json={'request_budget':1};s.commit()
        with operation(self.id,'plan_deck',model=True):record_usage(SimpleNamespace(usage={'input_tokens':12,'output_tokens':8}))
        with self.assertRaises(BudgetExhausted):
            with operation(self.id,'plan_design',model=True):self.fail('Must not execute')
        self.assertEqual(self.state().workflow.model_calls,1)
        self.assertEqual(self.state().workflow.input_tokens,12)
