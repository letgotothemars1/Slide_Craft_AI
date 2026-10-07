"""Notes are first-draft data, revision guarded, and native PowerPoint notes."""
import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from pptx import Presentation
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from fastapi import HTTPException
from app import repository
from app.db import Base
from app.project_schemas import ProjectCreateRequest, ProjectSlide, SpeakerNotesEditRequest
from app.routers.projects import edit_speaker_notes
from app.services.modular_build import build_slide_from_outline
from app.services.outline_service import starter_outline
from app.services.modular_export import render_project_pptx
from app.services.modular_slide_llm import _model_body
from app.services.llm_service import OpenAILLMService

class SpeakerNotesTest(unittest.TestCase):
    def setUp(self):
        self.variants_patch=patch('app.services.composition_choices.ensure_composition_variants');self.variants_patch.start()
        self.tmp=tempfile.TemporaryDirectory()
        self.engine=create_engine(f"sqlite:///{Path(self.tmp.name)/'notes.db'}")
        Base.metadata.create_all(self.engine)
        self.sessions=sessionmaker(bind=self.engine,expire_on_commit=False)
        with self.sessions() as s:
            p=repository.create_project(s,ProjectCreateRequest(assignment_text='Five slides',context_pack_text='Thesis: evidence then decide.'))
            outline=starter_outline(p.context_pack_text)
            p.outline_json=[i.model_dump() for i in outline]
            p.slides_json=[build_slide_from_outline(i).model_dump() for i in outline]
            p.phase='ready';s.commit();self.id=p.id
    def tearDown(self):
        self.variants_patch.stop()
        self.engine.dispose();self.tmp.cleanup()
    def state(self):
        with self.sessions() as s:return repository.project_response(s,repository.get_project(s,self.id))
    def test_stream_notes_same_call_after_content_and_legacy_response_compatible(self):
        service=OpenAILLMService.__new__(OpenAILLMService);service.client=MagicMock();service.model='mock'
        raw=json.dumps({'body':'A complete grounded observation.', 'sections':[{'heading':'Observation','text':'A complete grounded observation.'}], 'speaker_notes':'Explain the observation and its limitations. Then introduce the next step.'})
        service.client.responses.create.return_value=iter([SimpleNamespace(type='response.output_text.delta',delta=raw[:45]),SimpleNamespace(type='response.output_text.delta',delta=raw[45:]),SimpleNamespace(type='response.completed',response=None)])
        bodies,sections,notes=[],[],[]
        with patch('app.services.modular_slide_llm.get_llm_service',return_value=service):
            result=_model_body('s','u',False,on_partial=bodies.append,on_sections=sections.append,on_notes=notes.append)
        self.assertEqual(result,'A complete grounded observation.')
        self.assertTrue(bodies[0]);self.assertEqual(len(notes),1)
        schema=service.client.responses.create.call_args.kwargs['text']['format']['schema']
        self.assertEqual(schema['required'],['body','sections','speaker_notes'])
        self.assertEqual(service.client.responses.create.call_count,1)
        service.client.responses.create.return_value=SimpleNamespace(output_text='{"body":"Legacy body"}')
        service._extract_output_text=lambda r:r.output_text
        with patch('app.services.modular_slide_llm.get_llm_service',return_value=service):
            self.assertEqual(_model_body('s','u',False,on_notes=notes.append),'Legacy body')
        self.assertEqual(notes[-1],'')
    def test_edit_preserves_content_invalidates_design_and_guards_revision(self):
        before=self.state();slide=before.slides[0]
        with self.sessions() as s:
            after=edit_speaker_notes(self.id,slide.id,SpeakerNotesEditRequest(expected_revision=before.revision,speaker_notes='Explain the evidence limitation. Transition to the next slide.'),s)
        self.assertEqual(after.slides[0].blocks,slide.blocks)
        self.assertEqual(after.slides[0].sections,slide.sections)
        self.assertEqual(after.slides[1:],before.slides[1:])
        self.assertEqual(after.phase,'outline_draft')
        self.assertIsNone(after.workflow.deck_design)
        with self.sessions() as s:
            with self.assertRaises(HTTPException) as caught:
                edit_speaker_notes(self.id,slide.id,SpeakerNotesEditRequest(expected_revision=before.revision,speaker_notes='Stale'),s)
        self.assertEqual(caught.exception.status_code,409)
    def test_ready_body_and_notes_are_atomic_and_stale_generation_cannot_overwrite(self):
        from app.services import live_draft, modular_build
        with self.sessions() as session:
            project=repository.get_project(session,self.id)
            rows=project.slides_json
            rows[0]['blocks']['body'].update(status='generating',revision=7)
            rows[0]['status']='generating'
            project.slides_json=__import__('copy').deepcopy(rows)
            __import__('sqlalchemy').orm.attributes.flag_modified(project,'slides_json')
            session.commit()
        slide_id=self.state().slides[0].id
        with patch.object(modular_build,'SessionLocal',self.sessions):
            live_draft.write_body(self.id,slide_id,7,'Accepted grounded text','ready',speaker_notes='Accepted talk track')
            live_draft.write_body(self.id,slide_id,7,'Stale text','ready',speaker_notes='Stale notes')
        slide=self.state().slides[0]
        self.assertEqual(slide.blocks.body.text,'Accepted grounded text')
        self.assertEqual(slide.speaker_notes,'Accepted talk track')

    def test_initial_invalid_section_copy_keeps_generated_body_and_notes(self):
        from app.services import live_draft, modular_build
        with self.sessions() as session:
            project=repository.get_project(session,self.id)
            rows=__import__('copy').deepcopy(project.slides_json)
            rows[0]['status']='error';project.slides_json=rows
            project.phase='drafting';project.build_mode='model';session.commit()
        body='The pilot changed from 42% to 29%, across the same synthetic measurement period.'
        def generate(project,item,**kwargs):
            kwargs['on_partial'](body[:30])
            kwargs['on_sections']([{'heading':'Change','text':'The pilot improved.'}])
            kwargs['on_notes']('Describe the full change and the synthetic evidence limitation.')
            return body
        with patch.object(live_draft,'SessionLocal',self.sessions),patch.object(modular_build,'SessionLocal',self.sessions),patch.object(live_draft,'generate_slide_body',side_effect=generate):
            live_draft.run_live_draft(self.id,only_slide_id='s1')
        slide=self.state().slides[0]
        self.assertEqual(slide.status,'ready');self.assertEqual(slide.blocks.body.text,body)
        self.assertEqual(len(slide.sections),1);self.assertEqual(slide.sections[0].text,body)
        self.assertEqual(slide.speaker_notes,'Describe the full change and the synthetic evidence limitation.')

    def test_regenerated_body_invalid_grouping_keeps_new_notes_and_exact_body(self):
        from app.services import modular_regenerate, modular_build
        with self.sessions() as session:
            project=repository.get_project(session,self.id)
            rows=__import__('copy').deepcopy(project.slides_json)
            rows[0]['blocks']['body'].update(status='generating',revision=7)
            rows[0]['speaker_notes']='Old talk track'
            project.slides_json=rows;session.commit()
        body='A revised complete explanation of the accepted evidence.'
        def generate(project,slide_id,key,on_notes=None,on_sections=None):
            on_notes('New talk track explains revised evidence limitations.')
            on_sections([{'heading':'Evidence','text':'Changed wording is invalid.'}])
            return body
        with patch.object(modular_regenerate,'SessionLocal',self.sessions),patch.object(modular_build,'SessionLocal',self.sessions),patch.object(modular_regenerate,'generate_block_text',side_effect=generate):
            modular_regenerate.run_regeneration(self.id,'s1','body',7)
        slide=self.state().slides[0]
        self.assertEqual(slide.blocks.body.text,body);self.assertEqual(slide.sections[0].text,body)
        self.assertEqual(slide.speaker_notes,'New talk track explains revised evidence limitations.')
        self.assertEqual(slide.sections_status,'ready')

    def test_notes_export_with_visible_source_and_legacy_defaults(self):
        project=self.state();project.slides[0].speaker_notes='A grounded spoken explanation.'
        project.slides[0].show_source=True
        deck=Presentation(BytesIO(render_project_pptx(project)))
        self.assertIn('A grounded spoken explanation.',deck.slides[0].notes_slide.notes_text_frame.text)
        self.assertIn('Source:',deck.slides[0].notes_slide.notes_text_frame.text)
        self.assertNotIn('A grounded spoken explanation.', '\n'.join(sh.text for sh in deck.slides[0].shapes if sh.has_text_frame))
        raw=project.slides[1].model_dump();raw.pop('speaker_notes');raw.pop('composition_preference')
        self.assertEqual(ProjectSlide.model_validate(raw).speaker_notes,'')
        self.assertIsNone(ProjectSlide.model_validate(raw).composition_preference)

if __name__=='__main__':unittest.main()
