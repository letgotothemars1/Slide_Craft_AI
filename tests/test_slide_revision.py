import tempfile
import unittest
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from fastapi import BackgroundTasks, HTTPException
from pptx import Presentation
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from app import repository
from app.db import Base
from app.project_schemas import ProjectCreateRequest, SlideRegenerateRequest, SourceVisibilityRequest, BlockEditRequest
from app.routers.projects import revise_slide, source_visibility, edit_slide_block
from app.services import slide_revision, modular_build
from app.services.modular_build import build_slide_from_outline
from app.services.outline_service import starter_outline
from app.services.modular_export import render_project_pptx
from app.services.modular_slide_llm import generate_slide_body


class SlideRevisionTest(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.engine = create_engine(f"sqlite:///{Path(self.directory.name)/'test.db'}")
        Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine)
        self.patches = [patch.object(slide_revision,'SessionLocal',self.sessions), patch.object(modular_build,'SessionLocal',self.sessions), patch('app.routers.projects.get_llm_service',return_value=object())]
        for p in self.patches: p.start()
        with self.sessions() as session:
            project = repository.create_project(session,ProjectCreateRequest(assignment_text='Five slides',context_pack_text='Thesis: cautious'))
            self.id = project.id
            outline = starter_outline('Thesis: cautious')
            project.outline_json = [row.model_dump() for row in outline]
            project.slides_json = [build_slide_from_outline(row).model_dump() for row in outline]
            project.phase = 'ready'
            session.commit()

    def tearDown(self):
        for p in reversed(self.patches): p.stop()
        self.engine.dispose(); self.directory.cleanup()

    def state(self):
        with self.sessions() as session:
            return repository.project_response(session,repository.get_project(session,self.id))

    def start(self):
        state = self.state()
        background = BackgroundTasks()
        with self.sessions() as session:
            started = revise_slide(self.id,'s1',SlideRegenerateRequest(expected_revision=state.revision,instruction='Focus on limitations, use short steps.'),background,session)
        return {key:getattr(started.slides[0].blocks,key).revision for key in ('title','body')}

    def test_revision_stream_retains_other_slides_and_superseding_title(self):
        before = self.state()
        tokens = self.start()
        def generate(project,item,**kwargs):
            self.assertEqual(kwargs['instruction'],'Focus on limitations, use short steps.')
            kwargs['on_partial']('Live partial text')
            state = self.state()
            self.assertEqual(state.slides[0].blocks.body.text,'Live partial text')
            with self.sessions() as session:
                edit_slide_block(self.id,'s1','title',BlockEditRequest(expected_revision=state.revision,text='Accepted human title'),session)
            kwargs['on_title']('Generated title')
            return 'A revised cautious conclusion.'
        with patch.object(slide_revision,'generate_slide_body',side_effect=generate):
            slide_revision.run_slide_revision(self.id,'s1',tokens)
        after = self.state()
        self.assertEqual(after.slides[0].blocks.title.text,'Accepted human title')
        self.assertEqual(after.slides[0].blocks.body.text,'A revised cautious conclusion.')
        self.assertEqual(after.slides[1:],before.slides[1:])
        self.assertEqual(after.slides[0].blocks.source_label,before.slides[0].blocks.source_label)
        self.assertEqual(after.phase,'ready')
        self.assertEqual(after.slides[0].revision_instruction,'Focus on limitations, use short steps.')

    def test_failure_restores_previous_text_and_duplicate_is_rejected(self):
        before = self.state()
        tokens = self.start()
        with self.sessions() as session:
            with self.assertRaises(HTTPException) as raised:
                revise_slide(self.id,'s1',SlideRegenerateRequest(expected_revision=self.state().revision,instruction='Repeat'),BackgroundTasks(),session)
            self.assertEqual(raised.exception.status_code,409)
        def fail(project,item,**kwargs):
            kwargs['on_partial']('Unfinished response')
            raise ValueError('Deliberate model failure')
        with patch.object(slide_revision,'generate_slide_body',side_effect=fail), self.assertLogs(slide_revision.logger,level='ERROR'):
            slide_revision.run_slide_revision(self.id,'s1',tokens)
        after = self.state()
        for key in ('title','body'):
            self.assertEqual(getattr(after.slides[0].blocks,key).text,getattr(before.slides[0].blocks,key).text)
            self.assertEqual(getattr(after.slides[0].blocks,key).status,'error')
        self.assertEqual(after.slides[1:],before.slides[1:])

    def test_optional_footer_matches_pptx_and_hidden_source_is_in_notes(self):
        state = self.state()
        deck = Presentation(BytesIO(render_project_pptx(state)))
        texts = '\n'.join(shape.text for shape in deck.slides[0].shapes if shape.has_text_frame)
        self.assertNotIn('Source needed',texts)
        self.assertIn('Source needed',deck.slides[0].notes_slide.notes_text_frame.text)
        with self.sessions() as session:
            updated = source_visibility(self.id,'s1',SourceVisibilityRequest(expected_revision=state.revision,show_source=True),session)
        deck = Presentation(BytesIO(render_project_pptx(updated)))
        self.assertIn('Source needed','\n'.join(shape.text for shape in deck.slides[0].shapes if shape.has_text_frame))
        self.assertFalse(updated.slides[1].show_source)

    def test_instruction_is_passed_to_provider_prompt_and_title_is_requested(self):
        with self.sessions() as session:
            project = repository.get_project(session,self.id)
            item = starter_outline('Thesis: cautious')[0]
            titles=[]
            with patch('app.services.modular_slide_llm._model_body',return_value='Short revised body') as model:
                generate_slide_body(project,item,instruction='Focus on limitations.',on_partial=lambda _:None,on_title=titles.append)
            system,user,comparison = model.call_args.args
            self.assertIn('STUDENT REVISION REQUEST:\nFocus on limitations.',user)
            self.assertIn('Rewrite both the title JSON field',system)
            self.assertIs(model.call_args.kwargs['on_title'].__self__,titles)
