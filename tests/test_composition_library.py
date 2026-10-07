import unittest
from io import BytesIO
from pathlib import Path
from pptx import Presentation

from app.project_schemas import DesignPlan, DraftVisual, ProjectResponse, ProjectSlide, SlideBlock, SlideBlocks, SlideSection
from app.services.composition_catalog import COMPOSITION_CATALOG
from app.services.design_quality import render_scene
from app.services.design_scene import build_scene
from app.services.modular_export import render_project_pptx


def sample_slide(composition='balanced', arrangement='columns', visual=None):
    sections = [
        SlideSection(id='review', heading='Human review', text='Staff review takes 18 minutes per day.'),
        SlideSection(id='maintenance', heading='Maintenance', text='Camera maintenance takes two hours per week.'),
        SlideSection(id='scope', heading='Evaluation scope', text='The campus sample is narrow. Transfer to another site is unknown.'),
    ]
    block = lambda value: SlideBlock(text=value, status='ready', revision=1)
    return ProjectSlide(id='s3', status='ready', revision=1, sections=sections, show_source=True, visual=visual,
        blocks=SlideBlocks(title=block('Operating trade-offs'), body=block('\n\n'.join(s.text for s in sections)), source_label=block('synthetic.pdf, p. 2')),
        design=DesignPlan(layout='process' if visual and visual.kind=='process' else 'chart' if visual else 'editorial',
            emphasis='quiet', rationale='Preserve accepted sections.', visual=visual, composition=composition,
            arrangement=arrangement, focal_section_id='maintenance'))


class CompositionLibraryTest(unittest.TestCase):
    def assert_usable(self, slide, scene):
        self.assertEqual([e.text for e in scene if e.section_field=='text'], [s.text for s in slide.sections])
        self.assertEqual([e.text for e in scene if e.section_field=='heading'], [s.heading for s in slide.sections])
        self.assertEqual({e.section_id for e in scene if e.section_id}, {s.id for s in slide.sections})
        self.assertEqual(next(e.text for e in scene if e.block_key=='source_label'), 'synthetic.pdf, p. 2')
        for element in scene:
            self.assertGreaterEqual(element.x, 0)
            self.assertGreaterEqual(element.y, 0)
            self.assertGreater(element.w, 0)
            self.assertGreater(element.h, 0)
            self.assertLessEqual(element.x+element.w, 100.01)
            self.assertLessEqual(element.y+element.h, 56.26)
        _, issues = render_scene(scene)
        self.assertEqual(issues, [])

    def test_poster_keeps_long_title_words_intact(self):
        slide=sample_slide('poster')
        slide.blocks.title.text='Pilot observations'
        scene=build_scene(slide,'clean_editorial',1)
        title=next(e for e in scene if e.block_key=='title')
        # Native title width must fit the entire longest word, avoiding a
        # chopped mid-word title even when the box is tall enough.
        self.assertLessEqual(len('observations')*title.size*.64,title.w+.1)
        self.assertGreaterEqual(title.size,2.0)
        self.assertEqual(title.text,'Pilot observations')

    def test_catalog_has_distinct_executable_native_geometry(self):
        geometries = []
        for composition in COMPOSITION_CATALOG:
            slide = sample_slide(composition)
            scene = build_scene(slide, 'clean_editorial', 3)
            self.assert_usable(slide, scene)
            geometries.append([(e.x,e.y,e.w,e.h) for e in scene if e.section_field=='text'])
        self.assertEqual(len({str(g) for g in geometries}), 4)

    def test_choices_are_distinct_for_single_section_and_legacy_body(self):
        from app.services.composition_choices import get_composition_choices
        from app.project_schemas import OutlineItem
        from types import SimpleNamespace
        item=OutlineItem(id='s3',order=1,purpose='Explain',title='Operating trade-offs',key_message='Review matters',evidence_refs=[],layout_type='title')
        project=SimpleNamespace(theme='clean_editorial')
        for sections in ([sample_slide().sections[0]],[]):
            slide=sample_slide().model_copy(update={'sections':sections})
            choices=get_composition_choices(project,item,slide)
            self.assertEqual([c['id'] for c in choices],['balanced','bands','poster'])
            self.assertEqual([c['unconventional'] for c in choices],[False,False,True])
            self.assertEqual(len({str(c['scene']) for c in choices}),3)
            for choice in choices:
                scene=[__import__('app.project_schemas',fromlist=['SceneElement']).SceneElement.model_validate(e) for e in choice['scene']]
                _,issues=render_scene(scene)
                self.assertEqual(issues,[])
                self.assertIn(slide.blocks.title.text,[e.text for e in scene])
                if sections:
                    self.assertEqual([e.text for e in scene if e.section_field=='text'],[sections[0].text])
                else:
                    self.assertIn(slide.blocks.body.text,[e.text for e in scene])

    def test_choices_keep_visual_data_and_do_not_mutate_accepted_slide(self):
        from app.services.composition_choices import get_composition_choices
        from app.project_schemas import OutlineItem
        from types import SimpleNamespace
        item=OutlineItem(id='s3',order=3,purpose='Explain',title='Trade-offs',key_message='Review matters',evidence_refs=[],layout_type='content')
        for visual in [DraftVisual(kind='bars',labels=['Before','After'],values=[42,29],unit='%'),DraftVisual(kind='process',labels=['Detect','Review','Improve'],values=[],unit='')]:
            slide=sample_slide(visual=visual)
            before=slide.model_dump()
            for theme in ('clean_editorial','dark_tech_pitch','infographic_bright'):
                choices=get_composition_choices(SimpleNamespace(theme=theme),item,slide)
                self.assertEqual(len({str(c['scene']) for c in choices}),3)
                for choice in choices:
                    self.assertEqual(choice['design']['visual'],visual.model_dump())
                    self.assertEqual([e['text'] for e in choice['scene'] if e['section_field']=='text'],[s.text for s in slide.sections])
            self.assertEqual(slide.model_dump(),before)

    def test_grounded_callouts_preserve_full_claims_and_qualifications(self):
        from app.services.design_scene import _grounded_callouts
        slide=sample_slide('feature')
        slide.sections=[SlideSection(id='observation',heading='Observed change',text='In the synthetic case, the share changed from 42% to 29% (p. 1).')]
        scene=build_scene(slide,'clean_editorial',3)
        self.assertEqual(_grounded_callouts(slide.sections),['42%','29%'])
        self.assertEqual([e.text for e in scene if e.section_field=='text'],[slide.sections[0].text])
        self.assertTrue(all(next(e for e in scene if e.text==value).size>=6 for value in ['42%','29%']))
        self.assertNotIn('13%', [e.text for e in scene])
        slide.sections=[SlideSection(id='review',heading='Staff review',text='Staff review: 18 minutes per day'),SlideSection(id='maintenance',heading='Maintenance',text='Camera maintenance: two hours per week')]
        scene=build_scene(slide,'clean_editorial',4)
        self.assertEqual(_grounded_callouts(slide.sections),['18 minutes per day','two hours per week'])
        self.assertEqual([e.text for e in scene if e.section_field=='text'],[s.text for s in slide.sections])
        self.assertIn('two',[e.text for e in scene]);self.assertNotIn('2',[e.text for e in scene])
        self.assertEqual(_grounded_callouts([SlideSection(id='none',heading='Context',text='The team discussed 42% and 29%, but no transition was established.')]),[])
        _,issues=render_scene(scene);self.assertEqual(issues,[])

    def test_feature_allocates_width_to_dense_support_and_compact_hero(self):
        slide=sample_slide('feature')
        slide.sections=[SlideSection(id='takeaway',heading='Takeaway',text='Promising, not proven.'),SlideSection(id='limits',heading='Limits',text='This synthetic pilot has no randomized control group. The study cannot establish causation. A broader evaluation is required before making deployment claims. Treat the result as a pilot signal and assess transfer to other sites.')]
        slide.design.focal_section_id='takeaway'
        scene=build_scene(slide,'clean_editorial',5)
        bodies={e.section_id:e for e in scene if e.section_field=='text'}
        self.assertGreater(bodies['limits'].w,bodies['takeaway'].w)
        slide.design.composition='poster';slide.design.layout='hero'
        slide.blocks.title.text='North Campus sorting pilot: why it matters'
        slide.sections=[SlideSection(id='why',heading='Why it matters',text='Sorting errors make it harder to keep reusable materials separate on a campus.')]
        scene=build_scene(slide,'clean_editorial',1)
        title=next(e for e in scene if e.block_key=='title')
        self.assertGreater(title.w,35)
        _,issues=render_scene(scene);self.assertEqual(issues,[])

    def test_feature_emphasizes_requested_section_and_keeps_binding_order(self):
        slide = sample_slide('feature')
        scene = build_scene(slide, 'clean_editorial', 3)
        bodies = {e.section_id:e for e in scene if e.section_field=='text'}
        self.assertGreaterEqual(bodies['maintenance'].w, 32)
        self.assertGreater(bodies['maintenance'].size, bodies['review'].size)
        self.assert_usable(slide, scene)

    def test_draft_scene_is_neutral_independent_of_final_theme(self):
        slide = sample_slide().model_copy(update={'design':None})
        scenes = [build_scene(slide, theme, 3) for theme in ('clean_editorial','dark_tech_pitch','infographic_bright')]
        self.assertEqual(scenes[0], scenes[1])
        self.assertEqual(scenes[0], scenes[2])
        self.assertEqual(scenes[0][0].color, '#FFFFFF')

    def test_process_and_chart_keep_accepted_visuals_editable(self):
        visuals = [DraftVisual(kind='process', labels=['Detect','Staff review','Improve'], values=[], unit=''),
                   DraftVisual(kind='bars', labels=['Before','After'], values=[42,29], unit='%')]
        for visual in visuals:
            for composition in COMPOSITION_CATALOG:
                slide = sample_slide(composition, visual=visual)
                scene = build_scene(slide, 'clean_editorial', 3)
                self.assert_usable(slide, scene)
                texts = [e.text for e in scene if e.kind=='text']
                for label in visual.labels:
                    self.assertIn(label, texts)
                for value in visual.values:
                    self.assertIn(f'{value:g}{visual.unit}', texts)

    def test_four_sections_and_longer_content_remain_readable(self):
        slide = sample_slide('bands')
        slide.sections.append(SlideSection(id='decision', heading='Decision', text='Keep the pilot narrow and assess the workload before scaling.'))
        slide.sections[2].text = 'The campus sample is narrow. Transfer to another site is unknown. A broader evaluation is required before making deployment claims.'
        scene = build_scene(slide, 'clean_editorial', 3)
        self.assert_usable(slide, scene)

    def test_export_uses_scene_geometry_and_native_editable_text(self):
        fixture=Path(__file__).resolve().parents[1]/'project-instructions'/'fixtures'/'demo-project.json'
        project=ProjectResponse.model_validate_json(fixture.read_text())
        project.phase='ready'
        project.theme='clean_editorial'
        project.slides=[]
        for index,item in enumerate(project.outline):
            item.layout_type='content'
            slide=sample_slide(list(COMPOSITION_CATALOG)[index%3])
            slide.id=item.id
            slide.design_status='ready'
            project.slides.append(slide)
        deck=Presentation(BytesIO(render_project_pptx(project)))
        for slide,item,pptx_slide in zip(project.slides,project.outline,deck.slides):
            scene=build_scene(slide,project.theme,item.order)
            native=[shape for shape in pptx_slide.shapes if shape.has_text_frame and shape.text]
            elements=[element for element in scene if element.kind=='text']
            self.assertEqual([shape.text for shape in native],[element.text for element in elements])
            self.assertFalse(any(shape.shape_type==13 for shape in pptx_slide.shapes))
            for shape,element in zip(native,elements):
                self.assertAlmostEqual(shape.left/914400,element.x*13.333/100,places=5)
                self.assertAlmostEqual(shape.top/914400,element.y*13.333/100,places=5)


if __name__ == '__main__':
    unittest.main()
