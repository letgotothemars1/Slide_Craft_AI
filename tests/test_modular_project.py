"""Revision and export checks for the key-free modular project path."""

import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from fastapi import BackgroundTasks, HTTPException
from pptx import Presentation
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app import repository
from app.db import Base
from app.project_schemas import (
    BlockEditRequest, OutlineApproveRequest, OutlineRevisionRequest, ProjectCreateRequest,
)
from app.routers.projects import (
    approve_outline, build_project, edit_slide_block, generate_outline, reset_slide_block,
)
from app.services import modular_build
from app.services.modular_export import render_project_pptx


class ModularProjectTest(unittest.TestCase):
    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        self.engine = create_engine(f"sqlite:///{Path(self.directory.name) / 'mvp.db'}")
        Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine)
        self.patcher = patch.object(modular_build, "SessionLocal", self.sessions)
        self.patcher.start()
        with self.sessions() as session:
            project = repository.create_project(session, ProjectCreateRequest(
                assignment_text="Make five slides", context_pack_text="Thesis: review matters.",
                theme="dark_tech_pitch",
            ))
            self.project_id = project.id
            draft = generate_outline(project.id, OutlineRevisionRequest(expected_revision=0), session)
            approved = approve_outline(project.id, OutlineApproveRequest(
                expected_revision=draft.revision, theme="dark_tech_pitch"), session)
            build_project(project.id, OutlineRevisionRequest(expected_revision=approved.revision), BackgroundTasks(), session)

    def tearDown(self) -> None:
        self.patcher.stop()
        self.engine.dispose()
        self.directory.cleanup()

    def state(self):
        with self.sessions() as session:
            return repository.project_response(session, repository.get_project(session, self.project_id))

    def test_progress_edits_conflict_and_export(self) -> None:
        snapshots = []

        def after_slide(slide_id: str) -> None:
            state = self.state()
            snapshots.append((slide_id, [slide.status for slide in state.slides]))
            if slide_id == "s1":
                with self.sessions() as session:
                    edit_slide_block(self.project_id, "s1", "body", BlockEditRequest(
                        expected_revision=state.revision, text="Reviewed body text"), session)

        modular_build.run_build(self.project_id, after_slide=after_slide)
        final = self.state()
        self.assertEqual(snapshots[0][1], ["ready", "queued", "queued", "queued", "queued"])
        self.assertEqual(final.phase, "ready")
        self.assertEqual(final.slides[0].blocks.body.text, "Reviewed body text")
        self.assertEqual(len(final.slides), 5)

        with self.sessions() as session:
            with self.assertRaises(HTTPException) as caught:
                edit_slide_block(self.project_id, "s1", "title", BlockEditRequest(
                    expected_revision=final.revision - 1, text="Stale title"), session)
            self.assertEqual(caught.exception.status_code, 409)

        with self.sessions() as session:
            edited = edit_slide_block(self.project_id, "s1", "title", BlockEditRequest(
                expected_revision=final.revision, text="Edited title"), session)
            reset = reset_slide_block(self.project_id, "s1", "title", OutlineRevisionRequest(
                expected_revision=edited.revision), session)
        self.assertEqual(reset.slides[0].blocks.title.text, reset.outline[0].title)
        self.assertEqual(reset.slides[0].blocks.body.text, "Reviewed body text")
        self.assertEqual(reset.slides[1], final.slides[1])

        deck = Presentation(BytesIO(render_project_pptx(reset)))
        self.assertEqual(len(deck.slides), 5)
        text = " ".join(shape.text for shape in deck.slides[0].shapes if shape.has_text_frame)
        self.assertIn("Reviewed body text", text)
        self.assertEqual(str(deck.slides[0].shapes[0].fill.fore_color.rgb), "0B1020")

    def test_one_slide_failure_and_retry(self) -> None:
        def failing_builder(item):
            if item.id == "s3":
                raise RuntimeError("deliberate failure")
            return modular_build.build_slide_from_outline(item)

        with patch.object(modular_build.logger, "exception"):
            modular_build.run_build(self.project_id, builder=failing_builder)
        failed = self.state()
        self.assertEqual(failed.phase, "error")
        self.assertEqual([slide.status for slide in failed.slides], ["ready", "ready", "error", "ready", "ready"])
        unchanged = [slide.revision for slide in failed.slides]
        with self.sessions() as session:
            project = repository.get_project(session, self.project_id)
            project.phase = "building"
            session.commit()
        modular_build.run_build(self.project_id, only_slide_id="s3")
        ready = self.state()
        self.assertEqual(ready.phase, "ready")
        self.assertEqual([slide.revision for slide in ready.slides[:2] + ready.slides[3:]],
                         unchanged[:2] + unchanged[3:])


if __name__ == "__main__":
    unittest.main()
